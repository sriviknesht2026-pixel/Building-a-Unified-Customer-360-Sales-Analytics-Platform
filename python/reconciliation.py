"""
ECRMDP reconciliation & validation report (Sprint 3).

Compares every layer of the pipeline and writes docs/ECRMDP_Validation_Report.xlsx:
  1. Raw file  -> Bronze staging table : row counts and column totals must match
  2. Raw file  -> Silver file          : rows removed by cleansing (reported, not failed)
  3. Silver    -> Warehouse            : row counts, column totals and business IDs must match
  4. Warehouse key integrity           : no empty or orphan surrogate keys

Run from the repo root after the master ETL job has finished:
  pip install pandas openpyxl psycopg2-binary
  python python/reconciliation.py

Database settings are read from config/db.env (the same file deployment/setup_env.py uses).
Exit code is 0 when every check passes and 1 when any check fails, so it can gate a pipeline.
"""
import datetime as dt
import pathlib
import sys

import pandas as pd

REPO = pathlib.Path(__file__).resolve().parent.parent
DATA = REPO / "datasets"
REPORT = REPO / "docs" / "ECRMDP_Validation_Report.xlsx"
TOL = 0.01  # tolerance for money/decimal totals (rounding to NUMERIC(…,2))

# source -> raw file, id column, bronze table, silver file, warehouse table, warehouse id column, measures [(source col, warehouse col)]
SOURCES = [
    ("Customer Master", "01_customer_master.csv", "customer_id", "customer_master_01", "customer_master_silver.csv",
     "dim_customer", "customer_id", [("income", "income"), ("age", "age")]),
    ("Lead Management", "02_lead_management.xlsx", "lead_id", "lead_management_02", "lead_management_silver.csv",
     "fact_leads", "lead_id", [("lead_score", "lead_score")]),
    ("Opportunity Management", "03_opportunity_management.json", "opportunity_id", "opportunity_management_03",
     "opportunity_management_silver.csv", "fact_opportunities", "opportunity_id", [("opportunity_value", "opportunity_value")]),
    ("Sales Pipeline", "04_sales_pipeline.xml", "deal_id", "sales_pipeline_04", "sales_pipeline_silver.csv",
     "fact_sales", "deal_id", [("close_value", "close_value")]),
    ("Marketing Campaigns", "05_marketing_campaign.csv", "campaign_id", "marketing_campaign_05", "marketing_campaign_silver.csv",
     "dim_campaign", "campaign_id", [("acquisition_cost", "acquisition_cost"), ("roi", "roi")]),
    ("Customer Support", "06_customer_support_tickets.csv", "ticket_id", "customer_support_tickets_06",
     "customer_support_tickets_silver.csv", "fact_support", "ticket_id", [("resolution_hours", "resolution_hours")]),
    ("Contact Center", "07_contact_center_logs.csv", "interaction_id", "contact_center_logs_07", "contact_center_logs_silver.csv",
     "fact_contact_center", "interaction_id", [("duration_minutes", "duration_minutes")]),
    ("Website", "08_website_registration.csv", "registration_id", "website_registration_08", "website_registration_silver.csv",
     "fact_website", "registration_id", []),
    ("Mobile App", "09_mobile_application.csv", "app_event_id", "mobile_application_09", "mobile_application_silver.csv",
     "fact_mobile", "app_event_id", [("session_duration_minutes", "session_duration_minutes")]),
    ("Social Media", "10_social_media_engagement.csv", "engagement_id", "social_media_engagement_10",
     "social_media_engagement_silver.csv", "fact_social_media", "engagement_id", [("impressions", "impressions"), ("clicks", "clicks")]),
]

# distinct-value dimensions: (dimension, column, silver file, silver column)
DISTINCT_DIMS = [("dim_product", "product_name", "opportunity_management_silver.csv", "product"),
                 ("dim_sales_agent", "sales_agent", "lead_management_silver.csv", "sales_agent")]

# required keys: (fact, key column, dimension, dimension key)
KEYS = [
    ("fact_leads", "customer_key", "dim_customer", "customer_key"), ("fact_leads", "sales_agent_key", "dim_sales_agent", "sales_agent_key"),
    ("fact_leads", "date_key", "dim_date", "date_key"),
    ("fact_opportunities", "customer_key", "dim_customer", "customer_key"), ("fact_opportunities", "sales_agent_key", "dim_sales_agent", "sales_agent_key"),
    ("fact_opportunities", "product_key", "dim_product", "product_key"), ("fact_opportunities", "date_key", "dim_date", "date_key"),
    ("fact_sales", "customer_key", "dim_customer", "customer_key"), ("fact_sales", "sales_agent_key", "dim_sales_agent", "sales_agent_key"),
    ("fact_sales", "product_key", "dim_product", "product_key"), ("fact_sales", "date_key", "dim_date", "date_key"),
    ("fact_support", "customer_key", "dim_customer", "customer_key"),
    ("fact_contact_center", "customer_key", "dim_customer", "customer_key"), ("fact_contact_center", "date_key", "dim_date", "date_key"),
    ("fact_website", "customer_key", "dim_customer", "customer_key"), ("fact_website", "date_key", "dim_date", "date_key"),
    ("fact_mobile", "customer_key", "dim_customer", "customer_key"), ("fact_mobile", "date_key", "dim_date", "date_key"),
    ("fact_social_media", "customer_key", "dim_customer", "customer_key"), ("fact_social_media", "campaign_key", "dim_campaign", "campaign_key"),
    ("fact_social_media", "date_key", "dim_date", "date_key"),
]


def load_env():
    env = REPO / "config" / "db.env"
    if not env.exists():
        sys.exit("Missing config/db.env - copy config/db.env.example and fill in your values.")
    cfg = {}
    for line in env.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            cfg[k.strip()] = v.strip()
    return cfg


def read_file(path):
    p = str(path)
    if p.endswith(".csv"):
        return pd.read_csv(path)
    if p.endswith(".xlsx"):
        return pd.read_excel(path)
    if p.endswith(".json"):
        return pd.read_json(path)
    if p.endswith(".xml"):
        return pd.read_xml(path, parser="etree")
    raise ValueError(p)


def total(df, col):
    return round(float(pd.to_numeric(df[col], errors="coerce").sum()), 2)


def main():
    try:
        import psycopg2
    except ImportError:
        sys.exit("psycopg2 is not installed. Run:  pip install psycopg2-binary")
    cfg = load_env()
    con = psycopg2.connect(host=cfg["DB_HOST"], port=cfg["DB_PORT"], dbname=cfg["DB_NAME"],
                           user=cfg["DB_USER"], password=cfg["DB_PASSWORD"])
    cur = con.cursor()

    def q(sql):
        cur.execute(sql)
        return cur.fetchone()[0]

    rows = []

    def check(layer, source, name, expected, actual, kind="exact"):
        if kind == "info":
            status = "INFO"
        elif kind == "money":
            status = "PASS" if expected is not None and actual is not None and abs(float(expected) - float(actual)) <= TOL else "FAIL"
        else:
            status = "PASS" if expected == actual else "FAIL"
        rows.append([layer, source, name, expected, actual, status])

    for name, raw_f, idc, bronze, silver_f, wh, wh_id, measures in SOURCES:
        raw = read_file(DATA / raw_f)
        silver = read_file(DATA / silver_f)

        # 1. raw -> bronze
        check("1 Raw -> Bronze", name, f"Row count ({raw_f} vs {bronze})", len(raw), q(f"SELECT COUNT(*) FROM {bronze}"))
        for src_col, _ in measures:
            check("1 Raw -> Bronze", name, f"Total of {src_col}", total(raw, src_col),
                  round(float(q(f"SELECT COALESCE(SUM({src_col}),0) FROM {bronze}")), 2), "money")

        # 2. raw -> silver (cleansing)
        removed = len(raw) - len(silver)
        check("2 Raw -> Silver", name, "Rows removed by cleansing (duplicates)", len(raw), f"{len(silver)} kept, {removed} removed", "info")

        # 3. silver -> warehouse
        check("3 Silver -> Warehouse", name, f"Row count ({silver_f} vs {wh})", len(silver), q(f"SELECT COUNT(*) FROM {wh}"))
        for src_col, wh_col in measures:
            check("3 Silver -> Warehouse", name, f"Total of {wh_col}", total(silver, src_col),
                  round(float(q(f"SELECT COALESCE(SUM({wh_col}),0) FROM {wh}")), 2), "money")
        cur.execute(f"SELECT {wh_id} FROM {wh}")
        wh_ids = {r[0] for r in cur.fetchall()}
        silver_ids = set(silver[idc].astype(str))
        check("3 Silver -> Warehouse", name, f"Silver {idc}s missing from {wh}", 0, len(silver_ids - wh_ids))
        check("3 Silver -> Warehouse", name, f"{wh} {wh_id}s not in silver file", 0, len(wh_ids - silver_ids))

    for dim, col, silver_f, s_col in DISTINCT_DIMS:
        silver = read_file(DATA / silver_f)
        check("3 Silver -> Warehouse", dim, f"Distinct {s_col} values vs {dim} rows", int(silver[s_col].nunique()),
              q(f"SELECT COUNT(*) FROM {dim}"))
    check("3 Silver -> Warehouse", "dim_date", "Calendar rows (2025-01-01 + 731 days)", 731, q("SELECT COUNT(*) FROM dim_date"))
    check("3 Silver -> Warehouse", "dim_date", "Duplicate date_keys", 0,
          q("SELECT COUNT(*) - COUNT(DISTINCT date_key) FROM dim_date"))

    # 4. key integrity
    for fact, key, dim, dkey in KEYS:
        check("4 Key integrity", fact, f"{key}: empty", 0, q(f"SELECT COUNT(*) FROM {fact} WHERE {key} IS NULL"))
        check("4 Key integrity", fact, f"{key}: orphan (no match in {dim})", 0,
              q(f"SELECT COUNT(*) FROM {fact} f WHERE f.{key} IS NOT NULL AND NOT EXISTS "
                f"(SELECT 1 FROM {dim} d WHERE d.{dkey} = f.{key})"))
    check("4 Key integrity", "fact_support", "date_key: empty (expected - source has no date)",
          "2000 (all)", f"{q('SELECT COUNT(*) FROM fact_support WHERE date_key IS NULL')} empty", "info")
    con.close()

    write_report(rows, cfg)
    failed = [r for r in rows if r[5] == "FAIL"]
    passed = sum(1 for r in rows if r[5] == "PASS")
    for r in failed:
        print(f"FAIL  {r[0]} | {r[1]} | {r[2]} | expected {r[3]} got {r[4]}")
    print(f"\n{passed} passed, {len(failed)} failed, {sum(1 for r in rows if r[5] == 'INFO')} info  ->  {REPORT.relative_to(REPO)}")
    sys.exit(1 if failed else 0)


def write_report(rows, cfg):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    F = "Arial"
    fills = {"PASS": "E2EFDA", "FAIL": "F8CBAD", "INFO": "DDEBF7"}
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s["A1"] = "ECRMDP Reconciliation & Validation Report"
    s["A1"].font = Font(name=F, bold=True, size=14)
    meta = [("Run at", dt.datetime.now().strftime("%Y-%m-%d %H:%M")), ("Database", f"{cfg['DB_NAME']} on {cfg['DB_HOST']}:{cfg['DB_PORT']}"),
            ("Generated by", "python/reconciliation.py"), ("Money tolerance", f"±{TOL}")]
    for i, (k, v) in enumerate(meta, 3):
        s.cell(i, 1, k).font = Font(name=F, bold=True)
        s.cell(i, 2, v).font = Font(name=F)
    s["A8"] = "Result"
    s["A8"].font = Font(name=F, bold=True, size=11)
    for i, st in enumerate(["PASS", "FAIL", "INFO"], 9):
        c = s.cell(i, 1, st)
        c.font = Font(name=F, bold=True)
        c.fill = PatternFill("solid", fgColor=fills[st])
        s.cell(i, 2, f"=COUNTIF(Checks!$F:$F,A{i})").font = Font(name=F)
    s["A12"] = "Overall"
    s["A12"].font = Font(name=F, bold=True)
    s["B12"] = '=IF(B10=0,"ALL CHECKS PASSED","CHECKS FAILED - see Checks sheet")'
    s["B12"].font = Font(name=F, bold=True)
    s["A14"] = "By layer"
    s["A14"].font = Font(name=F, bold=True, size=11)
    for j, h in enumerate(["Layer", "PASS", "FAIL", "INFO"], 1):
        s.cell(15, j, h).font = Font(name=F, bold=True)
    layers = sorted({r[0] for r in rows})
    for i, layer in enumerate(layers, 16):
        s.cell(i, 1, layer).font = Font(name=F)
        for j, st in enumerate(["PASS", "FAIL", "INFO"], 2):
            s.cell(i, j, f'=COUNTIFS(Checks!$A:$A,$A{i},Checks!$F:$F,"{st}")').font = Font(name=F)
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 44
    for col in "CD":
        s.column_dimensions[col].width = 8

    c = wb.create_sheet("Checks")
    c.append(["Layer", "Source / table", "Check", "Expected", "Actual", "Status"])
    for cell in c[1]:
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.border = border
    for r in rows:
        c.append(r)
    for row in c.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=F, size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[5].fill = PatternFill("solid", fgColor=fills[row[5].value])
    for col, w in zip("ABCDEF", [22, 24, 52, 16, 22, 9]):
        c.column_dimensions[col].width = w
    c.freeze_panes = "A2"
    c.auto_filter.ref = c.dimensions
    REPORT.parent.mkdir(exist_ok=True)
    wb.save(REPORT)


if __name__ == "__main__":
    main()
